import csv
import io
from datetime import date, datetime
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from clinicshield.ingest import read_csv
from clinicshield.input_limits import (
    MAX_CELL_CHARACTERS,
    MAX_COLUMNS,
    MAX_INVENTORY_ROWS,
    MAX_WORKBOOK_MEMBERS,
    MAX_WORKBOOK_SHEETS,
    MAX_WORKBOOK_UNCOMPRESSED_BYTES,
    InputError,
    validate_upload,
)
from clinicshield.privacy import validate_headers


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time().isoformat() != "00:00:00":
            raise InputError("Date cells must not contain a time.")
        text = value.date().isoformat()
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)

    if len(text) > MAX_CELL_CHARACTERS:
        raise InputError("A workbook value exceeds the length limit.")
    return text


def check_workbook_archive(data):
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()

            if len(members) > MAX_WORKBOOK_MEMBERS:
                raise InputError("The workbook contains too many components.")

            total_size = sum(member.file_size for member in members)
            if total_size > MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                raise InputError(
                    "The decompressed workbook exceeds the size limit."
                )

            names = [member.filename for member in members]
            if len(names) != len(set(names)):
                raise InputError("The workbook has duplicate components.")

            if any(member.flag_bits & 1 for member in members):
                raise InputError("Encrypted workbooks are not supported.")

    except BadZipFile:
        raise InputError("Provide a valid Excel .xlsx workbook.") from None


def read_workbook(data: bytes) -> dict:
    """Read inventory sheets in memory without saving uploaded content."""
    validate_upload(data)
    check_workbook_archive(data)

    workbook = None

    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=False,
            keep_links=False,
        )

        if len(workbook.sheetnames) > MAX_WORKBOOK_SHEETS:
            raise InputError("The workbook contains too many sheets.")

        # Check every sheet, including hidden sheets.
        for sheet in workbook.worksheets:
            if (
                sheet.max_column is not None
                and sheet.max_column > MAX_COLUMNS
            ):
                raise InputError("A workbook sheet contains too many columns.")

            if (
                sheet.max_row is not None
                and sheet.max_row > MAX_INVENTORY_ROWS + 1
            ):
                raise InputError("A workbook sheet contains too many rows.")

            first_row = next(sheet.iter_rows(), ())
            headers = [cell_text(cell.value) for cell in first_row]

            if len(headers) > MAX_COLUMNS:
                raise InputError("A workbook sheet contains too many columns.")

            try:
                validate_headers(headers)
            except ValueError as exc:
                raise InputError(str(exc)) from None

            if any(cell.data_type == "f" for cell in first_row):
                raise InputError("Column names must not contain formulas.")

        if not {"assets", "software"}.issubset(workbook.sheetnames):
            raise InputError(
                "The workbook needs assets and software sheets."
            )

        result = {"assets": [], "software": [], "services": []}

        for kind in result:
            if kind not in workbook.sheetnames:
                continue

            sheet = workbook[kind]
            stream = io.StringIO(newline="")
            writer = csv.writer(stream)
            has_data = False
            width = 0

            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                if row_number > MAX_INVENTORY_ROWS + 1:
                    raise InputError(
                        "A workbook sheet contains too many rows."
                    )

                if len(row) > MAX_COLUMNS:
                    raise InputError(
                        "A workbook sheet contains too many columns."
                    )

                if any(cell.data_type == "f" for cell in row):
                    raise InputError(
                        "Inventory cells must contain values, not formulas."
                    )

                if any(cell.data_type == "e" for cell in row):
                    raise InputError(
                        "Inventory cells must not contain Excel errors."
                    )

                values = [cell_text(cell.value) for cell in row]

                if row_number == 1:
                    while values and not values[-1].strip():
                        values.pop()
                    width = len(values)
                    writer.writerow(values)
                    continue

                if any(value.strip() for value in values[width:]):
                    raise InputError(
                        "Inventory data extends beyond the named columns."
                    )

                values = values[:width]
                values.extend([""] * (width - len(values)))

                if any(value.strip() for value in values):
                    has_data = True

                writer.writerow(values)

            if kind == "services" and not has_data:
                continue

            result[kind] = read_csv(
                stream.getvalue().encode("utf-8"), kind
            )

        return result

    except InputError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
        ParseError,
        DefusedXmlException,
        EOFError,
    ):
        raise InputError("Provide a valid Excel .xlsx workbook.") from None
    finally:
        if workbook is not None:
            workbook.close()