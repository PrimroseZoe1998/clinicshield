import io
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from clinicshield import ingest, workbook
from clinicshield.checklist_input import read_checklist
from clinicshield.input_limits import (
    MAX_JSON_BYTES,
    MAX_UPLOAD_BYTES,
    InputError,
    validate_upload,
)
from clinicshield.nmap_input import read_nmap


PROJECT = Path(__file__).resolve().parents[1]
TEMPLATE = PROJECT / "templates" / "clinicshield_inventory.xlsx"


def book_bytes(edit):
    book = load_workbook(TEMPLATE)
    stream = io.BytesIO()
    try:
        edit(book)
        book.save(stream)
        return stream.getvalue()
    finally:
        book.close()


def zip_bytes(members):
    stream = io.BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return stream.getvalue()


def test_accepts_exact_upload_limit():
    validate_upload(b"1234", max_bytes=4)


def test_rejects_one_byte_over_limit():
    with pytest.raises(InputError, match="size limit"):
        validate_upload(b"12345", max_bytes=4)


@pytest.mark.parametrize("value", ["text", None, bytearray(b"text")])
def test_rejects_non_bytes(value):
    with pytest.raises(InputError, match="as bytes"):
        validate_upload(value)


@pytest.mark.parametrize("kind", ["assets", "software", "services"])
def test_csv_checks_upload_limit_before_parsing(kind):
    with pytest.raises(InputError, match="size limit"):
        ingest.read_csv(b"x" * (MAX_UPLOAD_BYTES + 1), kind)


def test_nmap_checks_upload_limit_before_parsing():
    with pytest.raises(InputError, match="size limit"):
        read_nmap(
            b"x" * (MAX_UPLOAD_BYTES + 1),
            authorized=True,
        )


def test_workbook_checks_upload_limit_before_parsing():
    with pytest.raises(InputError, match="size limit"):
        workbook.read_workbook(b"x" * (MAX_UPLOAD_BYTES + 1))


def test_checklist_checks_json_limit():
    with pytest.raises(InputError, match="size limit"):
        read_checklist(b"x" * (MAX_JSON_BYTES + 1), b"{}")


def test_followups_check_json_limit():
    answers = {
        f"Q{number:02d}": "Yes" for number in range(1, 26)
    }
    with pytest.raises(InputError, match="size limit"):
        read_checklist(
            json.dumps(answers).encode("utf-8"),
            b"x" * (MAX_JSON_BYTES + 1),
        )


def test_csv_limits_columns(monkeypatch):
    monkeypatch.setattr(ingest, "MAX_COLUMNS", 2)
    with pytest.raises(InputError, match="too many columns"):
        ingest.read_csv(b"a,b,c\n", "services")


def test_csv_limits_header_length(monkeypatch):
    monkeypatch.setattr(ingest, "MAX_CELL_CHARACTERS", 4)
    with pytest.raises(InputError, match="column name"):
        ingest.read_csv(b"asset_id\n", "services")


def test_csv_limits_value_length(monkeypatch):
    monkeypatch.setattr(ingest, "MAX_CELL_CHARACTERS", 10)
    data = (
        b"asset_id,port,protocol,service,exposure\n"
        b"PRIVATE_SENTINEL,80,tcp,http,internal\n"
    )
    with pytest.raises(InputError) as error:
        ingest.read_csv(data, "services")
    assert "length limit" in str(error.value)
    assert "PRIVATE_SENTINEL" not in str(error.value)


def test_csv_limits_rows_including_blank_records(monkeypatch):
    monkeypatch.setattr(ingest, "MAX_INVENTORY_ROWS", 1)
    data = (
        b"asset_id,port,protocol,service,exposure\n"
        b"TEST-01,80,tcp,http,internal\n"
        b",,,,\n"
    )
    with pytest.raises(InputError, match="too many data rows"):
        ingest.read_csv(data, "services")


def test_workbook_limits_decompressed_size(monkeypatch):
    monkeypatch.setattr(
        workbook, "MAX_WORKBOOK_UNCOMPRESSED_BYTES", 100
    )
    data = zip_bytes({"large.xml": b"x" * 101})
    with pytest.raises(InputError, match="decompressed"):
        workbook.read_workbook(data)


def test_workbook_limits_archive_members(monkeypatch):
    monkeypatch.setattr(workbook, "MAX_WORKBOOK_MEMBERS", 1)
    data = zip_bytes({"one.xml": b"x", "two.xml": b"x"})
    with pytest.raises(InputError, match="too many components"):
        workbook.read_workbook(data)


def test_workbook_limits_sheet_count(monkeypatch):
    monkeypatch.setattr(workbook, "MAX_WORKBOOK_SHEETS", 3)
    with pytest.raises(InputError, match="too many sheets"):
        workbook.read_workbook(TEMPLATE.read_bytes())


def test_workbook_limits_hidden_sheet_columns():
    def edit(book):
        sheet = book.create_sheet("Hidden")
        sheet.cell(1, 51, "extra")
        sheet.sheet_state = "hidden"

    with pytest.raises(InputError, match="too many columns"):
        workbook.read_workbook(book_bytes(edit))


def test_workbook_limits_rows(monkeypatch):
    monkeypatch.setattr(workbook, "MAX_INVENTORY_ROWS", 2)
    with pytest.raises(InputError, match="too many rows"):
        workbook.read_workbook(TEMPLATE.read_bytes())


def test_workbook_limits_cell_length_safely():
    def edit(book):
        book["assets"]["A2"] = "PRIVATE_SENTINEL" + ("x" * 10_000)

    with pytest.raises(InputError) as error:
        workbook.read_workbook(book_bytes(edit))
    assert "length limit" in str(error.value)
    assert "PRIVATE_SENTINEL" not in str(error.value)


def test_workbook_rejects_malformed_internal_xml():
    stream = io.BytesIO()
    with ZipFile(TEMPLATE) as source:
        with ZipFile(stream, "w", compression=ZIP_DEFLATED) as target:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename == "xl/workbook.xml":
                    content = b"<broken"
                target.writestr(member.filename, content)

    with pytest.raises(InputError, match="valid Excel"):
        workbook.read_workbook(stream.getvalue())