import io
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from clinicshield.ingest import InputError
from clinicshield.workbook import read_workbook
from zipfile import ZipFile

TEMPLATE = (
    Path(__file__).resolve().parents[1]
    / "templates"
    / "clinicshield_inventory.xlsx"
)


def template_book():
    return load_workbook(TEMPLATE)


def workbook_bytes(book):
    stream = io.BytesIO()
    try:
        book.save(stream)
        return stream.getvalue()
    finally:
        book.close()


def column_number(sheet, name):
    return next(
        cell.column for cell in sheet[1] if cell.value == name
    )


def test_reads_template():
    result = read_workbook(TEMPLATE.read_bytes())
    assert {name: len(rows) for name, rows in result.items()} == {
        "assets": 3,
        "software": 3,
        "services": 3,
    }


def test_services_sheet_is_optional():
    book = template_book()
    del book["services"]
    result = read_workbook(workbook_bytes(book))
    assert result["services"] == []
    assert len(result["assets"]) == 3


def test_services_sheet_can_have_headers_only():
    book = template_book()
    sheet = book["services"]
    sheet.delete_rows(2, sheet.max_row)
    result = read_workbook(workbook_bytes(book))
    assert result["services"] == []


@pytest.mark.parametrize("sheet_name", ["assets", "software"])
def test_rejects_missing_required_sheet(sheet_name):
    book = template_book()
    del book[sheet_name]
    with pytest.raises(InputError, match="needs assets and software"):
        read_workbook(workbook_bytes(book))


@pytest.mark.parametrize("state", ["visible", "hidden", "veryHidden"])
def test_checks_personal_information_headers_on_extra_sheets(state):
    book = template_book()
    sheet = book.create_sheet("Extra")
    sheet.append(["patient_name"])
    sheet.append(["PRIVATE_SENTINEL"])
    sheet.sheet_state = state

    with pytest.raises(InputError) as error:
        read_workbook(workbook_bytes(book))

    assert "personal information" in str(error.value)
    assert "PRIVATE_SENTINEL" not in str(error.value)


def test_privacy_check_precedes_inventory_validation():
    book = template_book()
    book["assets"]["A2"] = ""
    sheet = book.create_sheet("Hidden")
    sheet.append(["medical_record"])
    sheet.sheet_state = "hidden"

    with pytest.raises(InputError, match="personal information"):
        read_workbook(workbook_bytes(book))


def test_accepts_excel_date_cells():
    book = template_book()
    sheet = book["assets"]
    column = column_number(sheet, "last_patched")
    sheet.cell(2, column, datetime(2026, 10, 14))
    result = read_workbook(workbook_bytes(book))
    assert result["assets"][0].last_patched.isoformat() == "2026-10-14"


def test_rejects_date_cells_with_time():
    book = template_book()
    sheet = book["assets"]
    column = column_number(sheet, "last_patched")
    sheet.cell(2, column, datetime(2026, 10, 14, 12, 30))
    with pytest.raises(InputError, match="must not contain a time"):
        read_workbook(workbook_bytes(book))


def test_accepts_numeric_port_cells():
    book = template_book()
    sheet = book["services"]
    column = column_number(sheet, "port")
    sheet.cell(2, column, 8080)
    result = read_workbook(workbook_bytes(book))
    assert result["services"][0].port == 8080


def test_ignores_trailing_formatting():
    book = template_book()
    book["assets"]["Z10"].number_format = "0.00"
    result = read_workbook(workbook_bytes(book))
    assert len(result["assets"]) == 3


def test_rejects_data_beyond_headers():
    book = template_book()
    sheet = book["assets"]
    sheet.cell(2, sheet.max_column + 1, "PRIVATE_SENTINEL")
    with pytest.raises(InputError) as error:
        read_workbook(workbook_bytes(book))
    assert "beyond the named columns" in str(error.value)
    assert "PRIVATE_SENTINEL" not in str(error.value)


@pytest.mark.parametrize("cell", ["A1", "A2"])
def test_rejects_inventory_formulas(cell):
    book = template_book()
    book["assets"][cell] = '="PRIVATE_SENTINEL"'
    with pytest.raises(InputError) as error:
        read_workbook(workbook_bytes(book))
    assert "formulas" in str(error.value)
    assert "PRIVATE_SENTINEL" not in str(error.value)


def test_rejects_excel_error_cells():
    book = template_book()
    book["assets"]["A2"] = "#DIV/0!"
    with pytest.raises(InputError, match="Excel errors"):
        read_workbook(workbook_bytes(book))


def test_rejects_invalid_inventory_values_safely():
    book = template_book()
    sheet = book["services"]
    column = column_number(sheet, "port")
    sheet.cell(2, column, "PRIVATE_SENTINEL")
    with pytest.raises(InputError) as error:
        read_workbook(workbook_bytes(book))
    assert "invalid fields" in str(error.value)
    assert "PRIVATE_SENTINEL" not in str(error.value)


def test_rejects_empty_assets_sheet():
    book = template_book()
    sheet = book["assets"]
    sheet.delete_rows(2, sheet.max_row)
    with pytest.raises(InputError, match="at least one data row"):
        read_workbook(workbook_bytes(book))


def test_rejects_duplicate_assets():
    book = template_book()
    sheet = book["assets"]
    column = column_number(sheet, "asset_id")
    sheet.cell(3, column, sheet.cell(2, column).value)
    with pytest.raises(InputError, match="Asset IDs must be unique"):
        read_workbook(workbook_bytes(book))


@pytest.mark.parametrize("data", [b"", b"not an Excel workbook"])
def test_rejects_invalid_workbook(data):
    with pytest.raises(InputError, match="valid Excel"):
        read_workbook(data)


def test_rejects_non_bytes_input():
    with pytest.raises(InputError, match="as bytes"):
        read_workbook("not bytes")

def test_rejects_truncated_workbook():
    data = TEMPLATE.read_bytes()

    with pytest.raises(InputError, match="valid Excel"):
        read_workbook(data[:len(data) // 2])


def test_rejects_zip_without_excel_components():
    stream = io.BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("example.txt", "PRIVATE_SENTINEL")

    with pytest.raises(InputError, match="valid Excel") as error:
        read_workbook(stream.getvalue())

    assert "PRIVATE_SENTINEL" not in str(error.value)        